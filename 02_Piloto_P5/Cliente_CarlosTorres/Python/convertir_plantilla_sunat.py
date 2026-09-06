"""Convierte las plantillas oficiales de SUNAT (.xls legado) a .xlsx,
preservando el texto de encabezado y las celdas combinadas."""
import xlrd
import openpyxl
from pathlib import Path


def convertir_si_falta(xls_path: Path, xlsx_path: Path) -> Path:
    """Convierte xls_path a xlsx_path si este ultimo no existe todavia.
    Idempotente: si ya existe, no hace nada y retorna la ruta existente."""
    if xlsx_path.exists():
        return xlsx_path

    xlsx_path.parent.mkdir(parents=True, exist_ok=True)

    wb_origen = xlrd.open_workbook(str(xls_path), formatting_info=True)
    hoja_origen = wb_origen.sheet_by_index(0)

    wb_destino = openpyxl.Workbook()
    ws_destino = wb_destino.active
    ws_destino.title = hoja_origen.name.strip()[:31]

    for r in range(hoja_origen.nrows):
        for c in range(hoja_origen.ncols):
            valor = hoja_origen.cell_value(r, c)
            if valor != '':
                ws_destino.cell(row=r + 1, column=c + 1, value=valor)

    for rlo, rhi, clo, chi in hoja_origen.merged_cells:
        ws_destino.merge_cells(start_row=rlo + 1, end_row=rhi,
                                start_column=clo + 1, end_column=chi)

    wb_destino.save(str(xlsx_path))
    return xlsx_path
