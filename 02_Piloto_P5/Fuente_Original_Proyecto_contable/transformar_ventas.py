import pandas as pd
import numpy as np
from openpyxl import load_workbook

ARCHIVO = 'Data/IA_Ventas.xlsx'
HOJA_SALIDA = 'Asientos_Ventas'

DETALLE_CXC = 'EMITIDAS EN CARTERA - Facturas, Boletas y Otros Comprobantes Por Cobrar - Terceros'
DETALLE_VENTAS = 'TERCEROS - Prestacion De Servicios - Ventas'
DETALLE_IGV = 'IGV - CUENTA PROPIA - Por Pagar - Gobierno Central'
DETALLE_DETRAC = 'BANCO DE LA NACION CTA. CTE. S/. 00005-269121 DETRACCIONES'

PERIODO_MAP = {
    'Ene.25': 'En.25', 'Feb.25': 'Feb.25', 'Mar.25': 'Mar.25',
    'Abr.25': 'Abr.25', 'May.25': 'May.25', 'Jun.25': 'Jun.25',
    'Jul.25': 'Jul.25', 'Ago.25': 'Ago.25', 'Set.25': 'Set.25',
    'Oct.25': 'Oct.25', 'Nov.25': 'Nov.25', 'Dic.25': 'Dic.25',
}

def periodo_glosa(periodo):
    return PERIODO_MAP.get(str(periodo), str(periodo))

def safe(val):
    if pd.isna(val):
        return None
    return val

df = pd.read_excel(ARCHIVO, sheet_name='ventas')
print(f'Fuente: {len(df)} filas | Tipos CP: {df["Tipo CP/Doc."].unique()}')
print(f'Columnas disponibles: {df.columns.tolist()}')

# Identificar notas de crédito: Tipo=7 o IGV negativo (Total CP puede ser NaN si era fórmula)
df['_es_nc'] = (df['Tipo CP/Doc.'] == 7) | (df['IGV / IPM'] < 0)

journal_rows = []
nro = 1  # NRO correlativo por asiento

for i, row in df.iterrows():
    periodo = row['Periodo']
    fecha = row['Fecha de emisión']
    tipo = row['Tipo CP/Doc.']
    serie = str(row['Serie del CDP'])
    nro_cp = row['Nro CP o Doc. Nro Inicial (Rango)']
    cliente = str(row['Apellidos Nombres/ Razón Social'])
    igv = safe(row['IGV / IPM'])
    total = safe(row['Total CP'])
    detrac = safe(row['detrac'])
    pago_det = safe(row['Pago det'])
    es_nc = row['_es_nc']

    # BI = Valor Facturado (columna principal)
    valor_fact = safe(row['Valor Facturado Exportación'])
    bi = valor_fact

    # Derivar Total CP si es NaN (columna puede tener fórmulas que openpyxl no evalúa)
    if total is None and valor_fact is not None and igv is not None:
        total = round(abs(valor_fact) + abs(igv), 2)
        if es_nc:
            total = -total  # respetar signo de NC

    if igv is None:
        print(f'  [!] Fila {i} sin IGV — omitida')
        continue
    if total is None:
        print(f'  [!] Fila {i} sin Total (ni derivable) — omitida')
        continue

    per_label = periodo_glosa(periodo)
    nro_doc = f"{serie}-{nro_cp}"
    glosa_vta = f"Registro de Ventas {per_label}: {nro_doc} {cliente}"

    def fila(cuenta, detalle, debe, haber):
        return {
            'NRO': nro,
            'FECHA': fecha,
            'GLOSA': glosa_vta,
            'OBS': None,
            'CUENTA': cuenta,
            'DETALLE': detalle,
            'DEBE': round(debe, 2) if debe is not None else None,
            'HABER': round(haber, 2) if haber is not None else None,
            'CLIENTE/PROVEEDOR': cliente,
            'NRO. DOCUMENTO': nro_doc,
        }

    abs_total = abs(total)
    abs_bi = abs(bi) if bi is not None else None
    abs_igv = abs(igv)

    if es_nc:
        # Nota de crédito: reversa
        journal_rows += [
            fila(12121, DETALLE_CXC, None, abs_total),
            fila(7041, DETALLE_VENTAS, abs_bi, None),
            fila(40111, DETALLE_IGV, abs_igv, None),
        ]
    else:
        # Factura normal
        journal_rows += [
            fila(12121, DETALLE_CXC, abs_total, None),
            fila(7041, DETALLE_VENTAS, None, abs_bi),
            fila(40111, DETALLE_IGV, None, abs_igv),
        ]

    nro += 1

    # Asiento de pago de detracción (solo si hay fecha de pago)
    if detrac and detrac > 0 and pago_det is not None:
        # Validar fecha: si el año es anterior a 2020 es un typo en fuente
        pago_det_validado = pago_det
        if hasattr(pago_det, 'year') and pago_det.year < 2020:
            print(f"  [AVISO] Fecha de detracción anómala en {nro_doc}: {pago_det} — omitiendo asiento detracción")
            nro += 1
            continue
        glosa_det = f"Pago detraccion {per_label}: {nro_doc} {cliente}"

        def fila_det(cuenta, detalle, debe, haber):
            return {
                'NRO': nro,
                'FECHA': pago_det_validado,
                'GLOSA': glosa_det,
                'OBS': None,
                'CUENTA': cuenta,
                'DETALLE': detalle,
                'DEBE': round(debe, 2) if debe is not None else None,
                'HABER': round(haber, 2) if haber is not None else None,
                'CLIENTE/PROVEEDOR': cliente,
                'NRO. DOCUMENTO': nro_doc,
            }

        journal_rows += [
            fila_det(104201, DETALLE_DETRAC, detrac, None),
            fila_det(12121, DETALLE_CXC, None, detrac),
        ]
        nro += 1

df_journal = pd.DataFrame(journal_rows)

print(f'\n=== RESUMEN ASIENTOS ===')
print(f'Asientos generados: {nro - 1}')
print(f'Filas totales: {len(df_journal)}')
print(f'Suma DEBE: {df_journal["DEBE"].sum():.2f}')
print(f'Suma HABER: {df_journal["HABER"].sum():.2f}')
print(f'Diferencia (cuadre): {abs(df_journal["DEBE"].sum() - df_journal["HABER"].sum()):.2f}')
print()
print('Cuentas utilizadas:')
print(df_journal['CUENTA'].value_counts().sort_index())
print()
print('Primeras 15 filas:')
print(df_journal.head(15).to_string())

# Escribir en hoja nueva del Excel
try:
    wb = load_workbook(ARCHIVO)
    if HOJA_SALIDA in wb.sheetnames:
        del wb[HOJA_SALIDA]
    wb.save(ARCHIVO)
except Exception as e:
    print(f'[openpyxl load error] {e}')

# Fecha sin hora
df_journal['FECHA'] = pd.to_datetime(df_journal['FECHA']).dt.date

with pd.ExcelWriter(ARCHIVO, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
    df_journal.to_excel(writer, sheet_name=HOJA_SALIDA, index=False)
    ws = writer.sheets[HOJA_SALIDA]
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        for cell in row:
            cell.number_format = 'DD/MM/YYYY'

print(f'\nOK: Hoja "{HOJA_SALIDA}" escrita en {ARCHIVO}')
